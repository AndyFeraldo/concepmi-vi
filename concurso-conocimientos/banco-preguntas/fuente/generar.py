"""Genera el banco de preguntas del Concurso de Conocimientos (VI CONCEPMI).

Lee las preguntas de d1.py a d4.py y produce, en la carpeta superior:
  - banco-preguntas-capm.xlsx  (para revisar y editar)
  - preguntas.json             (para la aplicación)
  - preguntas.sql              (para cargar en Supabase desde el SQL Editor)

Uso:  python3 -I generar.py   (requiere openpyxl)
"""
import json
import random
import sys
from collections import Counter
from pathlib import Path

AQUI = Path(__file__).resolve().parent
sys.path.insert(0, str(AQUI))
from d1 import D1  # noqa: E402
from d2 import D2  # noqa: E402
from d3 import D3  # noqa: E402
from d4 import D4  # noqa: E402

from openpyxl import Workbook  # noqa: E402
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side  # noqa: E402
from openpyxl.worksheet.datavalidation import DataValidation  # noqa: E402

SALIDA = AQUI.parent
LETRAS = "ABCD"

DOMINIOS = [
    (1, "Fundamentos y conceptos básicos de la dirección de proyectos", "36 %", D1),
    (2, "Metodologías predictivas basadas en planes", "17 %", D2),
    (3, "Marcos y metodologías ágiles", "20 %", D3),
    (4, "Marcos de análisis de negocios", "27 %", D4),
]


def construir():
    preguntas = []
    for num, nombre, _, datos in DOMINIOS:
        for tema, dificultad, enunciado, correcta, distractores, explicacion in datos:
            preguntas.append(dict(dominio_num=num, dominio=nombre, tema=tema,
                                  dificultad=dificultad, enunciado=enunciado,
                                  correcta_txt=correcta, distractores=list(distractores),
                                  explicacion=explicacion))

    # Posición de la respuesta correcta: equilibrada (45 por letra) y aleatoria
    rnd = random.Random(2026)
    posiciones = list(LETRAS) * (len(preguntas) // 4) + list(LETRAS)[: len(preguntas) % 4]
    rnd.shuffle(posiciones)

    for i, (p, letra) in enumerate(zip(preguntas, posiciones), start=1):
        opciones = p.pop("distractores")
        rnd.shuffle(opciones)
        opciones.insert(LETRAS.index(letra), p.pop("correcta_txt"))
        p["id"] = i
        p["opciones"] = dict(zip(LETRAS, opciones))
        p["correcta"] = letra
    return preguntas


def escribir_json(preguntas):
    datos = [dict(id=p["id"], dominio=p["dominio"], tema=p["tema"], dificultad=p["dificultad"],
                  enunciado=p["enunciado"], opciones=p["opciones"], correcta=p["correcta"],
                  explicacion=p["explicacion"]) for p in preguntas]
    (SALIDA / "preguntas.json").write_text(json.dumps(datos, ensure_ascii=False, indent=2), encoding="utf-8")


def sql_txt(t):
    return "'" + t.replace("'", "''") + "'"


def escribir_sql(preguntas):
    filas = []
    for p in preguntas:
        o = p["opciones"]
        filas.append("(" + ", ".join([str(p["id"]), sql_txt(p["dominio"]), sql_txt(p["tema"]),
                                       sql_txt(p["dificultad"]), sql_txt(p["enunciado"]),
                                       sql_txt(o["A"]), sql_txt(o["B"]), sql_txt(o["C"]), sql_txt(o["D"]),
                                       sql_txt(p["correcta"]), sql_txt(p["explicacion"])]) + ")")
    sql = f"""-- Banco de preguntas · Concurso de Conocimientos · VI CONCEPMI
-- Pegar completo en Supabase > SQL Editor > Run (después de supabase/esquema.sql).
-- Se puede ejecutar varias veces: reemplaza las preguntas existentes con el mismo id.

create table if not exists public.preguntas (
  id           integer primary key,
  dominio      text not null,
  tema         text not null,
  dificultad   text not null,
  enunciado    text not null,
  opcion_a     text not null,
  opcion_b     text not null,
  opcion_c     text not null,
  opcion_d     text not null,
  correcta     char(1) not null check (correcta in ('A','B','C','D')),
  explicacion  text,
  activa       boolean not null default true
);

-- Privada: nadie lee las preguntas directamente (la respuesta correcta no sale de la base).
alter table public.preguntas enable row level security;
revoke all on public.preguntas from anon, authenticated;

insert into public.preguntas
  (id, dominio, tema, dificultad, enunciado, opcion_a, opcion_b, opcion_c, opcion_d, correcta, explicacion)
values
{(","+chr(10)).join(filas)}
on conflict (id) do update set
  dominio = excluded.dominio, tema = excluded.tema, dificultad = excluded.dificultad,
  enunciado = excluded.enunciado, opcion_a = excluded.opcion_a, opcion_b = excluded.opcion_b,
  opcion_c = excluded.opcion_c, opcion_d = excluded.opcion_d, correcta = excluded.correcta,
  explicacion = excluded.explicacion;
"""
    (SALIDA / "preguntas.sql").write_text(sql, encoding="utf-8")


# Paleta de la landing CONCEPMI VI
NOCHE, INDIGO, MORADO, LILA, CELESTE, NARANJA, NIEBLA = "170D56", "24107A", "4A17B5", "A98BFF", "1FCBEA", "FF6A1F", "E3DFED"
COLOR_DOMINIO = {1: "EDE7FF", 2: "E3F8FC", 3: "FFEDE3", 4: "E8FBF1"}


def escribir_xlsx(preguntas):
    wb = Workbook()
    borde = Border(*(Side(style="thin", color="D4CFE3"),) * 4)
    cab_fill = PatternFill("solid", fgColor=INDIGO)
    cab_font = Font(name="Poppins", bold=True, color="FFFFFF", size=10)
    envolver = Alignment(wrap_text=True, vertical="top")

    # Hoja 1: Preguntas
    ws = wb.active
    ws.title = "Preguntas"
    ws.sheet_properties.tabColor = MORADO
    ws.merge_cells("A1:L1")
    ws["A1"] = "BANCO DE PREGUNTAS · CONCURSO DE CONOCIMIENTOS · VI CONCEPMI 2026 · ESTILO EXAMEN CAPM"
    ws["A1"].font = Font(name="Poppins", bold=True, size=14, color="FFFFFF")
    ws["A1"].fill = PatternFill("solid", fgColor=NOCHE)
    ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 30

    cab = ["ID", "Dominio", "Tema", "Dificultad", "Pregunta", "Alternativa A", "Alternativa B",
           "Alternativa C", "Alternativa D", "Respuesta correcta", "Explicación", "Activa"]
    anchos = [6, 26, 22, 12, 55, 34, 34, 34, 34, 12, 45, 9]
    for c, (t, w) in enumerate(zip(cab, anchos), start=1):
        celda = ws.cell(row=2, column=c, value=t)
        celda.fill, celda.font, celda.border = cab_fill, cab_font, borde
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[celda.column_letter].width = w
    ws.row_dimensions[2].height = 30

    for f, p in enumerate(preguntas, start=3):
        o = p["opciones"]
        valores = [p["id"], f"{p['dominio_num']}. {p['dominio']}", p["tema"], p["dificultad"],
                   p["enunciado"], o["A"], o["B"], o["C"], o["D"], p["correcta"], p["explicacion"], "Sí"]
        relleno = PatternFill("solid", fgColor=COLOR_DOMINIO[p["dominio_num"]])
        for c, v in enumerate(valores, start=1):
            celda = ws.cell(row=f, column=c, value=v)
            celda.alignment, celda.border = envolver, borde
            celda.font = Font(name="Poppins", size=9)
            if c in (1, 2, 3, 4):
                celda.fill = relleno
        # Resaltar la alternativa correcta
        col_ok = 6 + LETRAS.index(p["correcta"])
        ws.cell(row=f, column=col_ok).fill = PatternFill("solid", fgColor="D3F9E4")
        ws.cell(row=f, column=col_ok).font = Font(name="Poppins", size=9, bold=True, color="0B6B3A")
        rc = ws.cell(row=f, column=10)
        rc.alignment = Alignment(horizontal="center", vertical="top")
        rc.font = Font(name="Poppins", size=11, bold=True, color=MORADO)
        ws.cell(row=f, column=12).alignment = Alignment(horizontal="center", vertical="top")

    ultima = 2 + len(preguntas)
    dv = DataValidation(type="list", formula1='"A,B,C,D"', allow_blank=False)
    dv2 = DataValidation(type="list", formula1='"Básica,Intermedia,Avanzada"', allow_blank=False)
    dv3 = DataValidation(type="list", formula1='"Sí,No"', allow_blank=False)
    for v, rango in ((dv, f"J3:J{ultima}"), (dv2, f"D3:D{ultima}"), (dv3, f"L3:L{ultima}")):
        ws.add_data_validation(v)
        v.add(rango)
    ws.freeze_panes = "E3"
    ws.auto_filter.ref = f"A2:L{ultima}"

    # Hoja 2: Resumen
    rs = wb.create_sheet("Resumen")
    rs.sheet_properties.tabColor = CELESTE
    rs["A1"] = "Distribución según el Esquema de Contenido del Examen CAPM (PMI)"
    rs["A1"].font = Font(name="Poppins", bold=True, size=13, color=NOCHE)
    cab = ["Dominio", "Peso en el examen", "Preguntas", "Básica", "Intermedia", "Avanzada"]
    for c, t in enumerate(cab, start=1):
        celda = rs.cell(row=3, column=c, value=t)
        celda.fill, celda.font, celda.border = cab_fill, cab_font, borde
        celda.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for f, (num, nombre, peso, _) in enumerate(DOMINIOS, start=4):
        del_dom = [p for p in preguntas if p["dominio_num"] == num]
        cnt = Counter(p["dificultad"] for p in del_dom)
        for c, v in enumerate([f"{num}. {nombre}", peso, len(del_dom), cnt["Básica"], cnt["Intermedia"], cnt["Avanzada"]], start=1):
            celda = rs.cell(row=f, column=c, value=v)
            celda.border, celda.font = borde, Font(name="Poppins", size=10)
            celda.fill = PatternFill("solid", fgColor=COLOR_DOMINIO[num])
            if c > 1:
                celda.alignment = Alignment(horizontal="center")
    tot = 4 + len(DOMINIOS)
    rs.cell(row=tot, column=1, value="TOTAL").font = Font(name="Poppins", bold=True)
    for c, col in enumerate("CDEF", start=3):
        celda = rs.cell(row=tot, column=c, value=f"=SUM({col}4:{col}{tot - 1})")
        celda.font, celda.alignment = Font(name="Poppins", bold=True), Alignment(horizontal="center")
    rs.cell(row=tot, column=2, value="100 %").alignment = Alignment(horizontal="center")

    letras = Counter(p["correcta"] for p in preguntas)
    rs.cell(row=tot + 2, column=1, value="Respuestas correctas por letra (equilibradas)").font = Font(name="Poppins", bold=True, color=NOCHE)
    for i, l in enumerate(LETRAS):
        rs.cell(row=tot + 3 + i, column=1, value=f"Alternativa {l}")
        rs.cell(row=tot + 3 + i, column=2, value=letras[l]).alignment = Alignment(horizontal="center")

    rs.cell(row=tot + 8, column=1, value="Uso en el concurso").font = Font(name="Poppins", bold=True, color=NOCHE)
    notas = [
        "Cada ronda usa 5 preguntas aleatorias; cada enfrentamiento usa 15 (3 rondas).",
        "Las preguntas no se repiten durante el evento: 180 preguntas alcanzan para 12 enfrentamientos.",
        "Para desactivar una pregunta sin borrarla, cambiar la columna \"Activa\" a \"No\".",
        "Referencias: Guía del PMBOK (7.ª ed.), Guía de Prácticas Ágiles, Guía Scrum y Business Analysis for Practitioners (PMI).",
    ]
    for i, n in enumerate(notas):
        rs.cell(row=tot + 9 + i, column=1, value=f"• {n}").font = Font(name="Poppins", size=10)
    rs.column_dimensions["A"].width = 62
    for col in "BCDEF":
        rs.column_dimensions[col].width = 16

    wb.save(SALIDA / "banco-preguntas-capm.xlsx")


if __name__ == "__main__":
    preguntas = construir()
    escribir_json(preguntas)
    escribir_sql(preguntas)
    escribir_xlsx(preguntas)
    print(len(preguntas), "preguntas ·", dict(Counter(p["correcta"] for p in preguntas)))
